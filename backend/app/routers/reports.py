"""Played-game statistics and exports."""

import csv
import io

from fastapi import APIRouter, Depends, Query, Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..auth import current_user
from ..database import get_db
from ..game_service import accessible_session
from ..models import Answer, GameSession, User
from ..services import answer_distribution, leaderboard


router = APIRouter(prefix="/api/sessions", tags=["reports"])


@router.get("/{session_id}/statistics")
def session_statistics(session_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    game = accessible_session(db, session_id, user.id)
    board = leaderboard(db, game.id, update_ranks=False)
    answers = db.scalars(select(Answer).where(Answer.session_id == game.id)).all()
    by_question: dict[int, list[Answer]] = {}
    for answer in answers:
        by_question.setdefault(answer.question_id, []).append(answer)
    question_stats = []
    for question in game.quiz.questions:
        question_answers = by_question.get(question.id, [])
        correct_count = sum(1 for answer in question_answers if answer.is_correct)
        question_stats.append(
            {
                "id": question.id,
                "position": question.position,
                "text": question.text,
                "type": question.type,
                "options": question.options,
                "match_options": question.match_options or [],
                "correct_options": question.correct_options,
                "distribution": answer_distribution(db, game, question),
                "response_count": len(question_answers),
                "correct_count": correct_count,
                "correct_rate": round(correct_count / len(game.participants) * 100, 1) if game.participants else 0,
            }
        )
    total_possible = len(game.participants) * len(game.quiz.questions)
    total_correct = sum(1 for answer in answers if answer.is_correct)
    return {
        "session": {
            "id": game.id,
            "quiz_title": game.quiz.title,
            "course_tag": game.quiz.course_tag,
            "pin": game.pin,
            "status": game.status,
            "started_at": game.started_at,
            "ended_at": game.ended_at,
            "host": {"id": (game.host or game.quiz.owner).id, "name": (game.host or game.quiz.owner).name},
        },
        "summary": {
            "participants": len(game.participants),
            "questions": len(game.quiz.questions),
            "average_accuracy": round(total_correct / total_possible * 100, 1) if total_possible else 0,
            "average_score": round(sum(row["score"] for row in board) / len(board)) if board else 0,
        },
        "leaderboard": board,
        "questions": question_stats,
    }


def export_rows(db: Session, game: GameSession):
    board = leaderboard(db, game.id, update_ranks=False)
    rank_by_id = {row["participant_id"]: row for row in board}
    answers = db.scalars(select(Answer).where(Answer.session_id == game.id)).all()
    answer_map = {(answer.participant_id, answer.question_id): answer for answer in answers}
    rows = []
    for participant in game.participants:
        marks = []
        correct = 0
        for question in game.quiz.questions:
            answer = answer_map.get((participant.id, question.id))
            value = 1 if answer and answer.is_correct else 0
            marks.append(value)
            correct += value
        score = rank_by_id[participant.id]["score"]
        total = len(game.quiz.questions)
        rows.append(
            [
                participant.display_name,
                participant.student_id or "",
                *marks,
                correct,
                round(correct / total * 100, 2) if total else 0,
                score,
                rank_by_id[participant.id]["rank"],
            ]
        )
    rows.sort(key=lambda row: row[-1])
    return rows


@router.get("/{session_id}/export")
def export(
    session_id: int,
    format: str = Query("xlsx", pattern="^(xlsx|csv)$"),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    game = accessible_session(db, session_id, user.id)
    headers = [
        "Student Name",
        "Student ID",
        *[f"Q{index + 1}" for index in range(len(game.quiz.questions))],
        "Total Correct",
        "Accuracy %",
        "Total Score",
        "Rank",
    ]
    rows = export_rows(db, game)
    stem = f"{game.quiz.course_tag or 'quiz'}-{game.pin}-results"
    if format == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(headers)
        writer.writerows(rows)
        return Response(output.getvalue(), media_type="text/csv", headers={"Content-Disposition": f'attachment; filename="{stem}.csv"'})
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Results"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="16324F")
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 32)
    binary = io.BytesIO()
    workbook.save(binary)
    return Response(
        binary.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{stem}.xlsx"'},
    )
