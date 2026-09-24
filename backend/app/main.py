from __future__ import annotations

import asyncio
import csv
import io
import random
import secrets
from contextlib import asynccontextmanager
from datetime import timedelta

from fastapi import BackgroundTasks, Depends, FastAPI, HTTPException, Query, Response, WebSocket, WebSocketDisconnect
from fastapi.middleware.cors import CORSMiddleware
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import delete, func, inspect, or_, select, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, selectinload

from .auth import create_password_reset_token, create_token, current_user, decode_password_reset_token, decode_user_id, hash_password, password_fingerprint, verify_password
from .config import ALLOW_PUBLIC_REGISTRATION, CORS_ORIGINS, CORS_ORIGIN_REGEX, DEFAULT_INSTRUCTOR_EMAIL, DEFAULT_INSTRUCTOR_NAME, DEFAULT_INSTRUCTOR_PASSWORD, PASSWORD_RESET_BASE_URL, PASSWORD_RESET_ENABLED, PASSWORD_RESET_EXPIRE_MINUTES, RESULT_DISPLAY_SECONDS
from .database import Base, SessionLocal, engine, get_db
from .mailer import send_password_reset_email
from .models import Answer, CourseFolder, FolderMember, GameSession, Participant, Question, Quiz, User, utcnow
from .realtime import Client, manager
from .schemas import FolderIn, FolderMemberIn, ForgotPasswordIn, JoinIn, LoginIn, QuizIn, RegisterIn, ResetPasswordIn
from .services import answer_distribution, expected_submission, iso, leaderboard, question_payload, reveal_payload, session_snapshot


def initialize_database():
    Base.metadata.create_all(engine)
    schema = inspect(engine)
    quiz_columns = {column["name"] for column in schema.get_columns("quizzes")}
    question_columns = {column["name"] for column in schema.get_columns("questions")}
    session_columns = {column["name"] for column in schema.get_columns("sessions")}
    with engine.begin() as connection:
        if "folder_id" not in quiz_columns:
            connection.execute(text("ALTER TABLE quizzes ADD COLUMN folder_id INTEGER"))
        if "deleted_at" not in quiz_columns:
            connection.execute(text("ALTER TABLE quizzes ADD COLUMN deleted_at DATETIME"))
        if "host_id" not in session_columns:
            connection.execute(text("ALTER TABLE sessions ADD COLUMN host_id INTEGER"))
        if "match_options" not in question_columns:
            connection.execute(text("ALTER TABLE questions ADD COLUMN match_options JSON"))
        connection.execute(text("CREATE INDEX IF NOT EXISTS idx_quizzes_folder_id ON quizzes (folder_id)"))
        connection.execute(text("CREATE INDEX IF NOT EXISTS idx_quizzes_deleted_at ON quizzes (deleted_at)"))
        connection.execute(text("CREATE INDEX IF NOT EXISTS idx_sessions_host_id ON sessions (host_id)"))
        connection.execute(text("UPDATE sessions SET host_id = (SELECT owner_id FROM quizzes WHERE quizzes.id = sessions.quiz_id) WHERE host_id IS NULL"))
        if engine.dialect.name == "sqlite":
            connection.execute(text("PRAGMA optimize"))
    with SessionLocal() as db:
        if not db.scalar(select(User).where(User.email == DEFAULT_INSTRUCTOR_EMAIL)):
            db.add(User(name=DEFAULT_INSTRUCTOR_NAME, email=DEFAULT_INSTRUCTOR_EMAIL, password_hash=hash_password(DEFAULT_INSTRUCTOR_PASSWORD)))
            db.commit()


@asynccontextmanager
async def lifespan(_app: FastAPI):
    initialize_database()
    with SessionLocal() as db:
        active = db.scalars(
            select(GameSession)
            .where(GameSession.status == "live")
            .options(selectinload(GameSession.quiz).selectinload(Quiz.questions))
        ).all()
        for game in active:
            if game.current_question_index is None:
                continue
            question = game.quiz.questions[game.current_question_index]
            if game.is_revealed:
                schedule_advance(game.id, question.id)
            else:
                remaining = max(0.0, ((game.question_deadline_at or utcnow()) - utcnow()).total_seconds())
                schedule_reveal(game.id, question.id, remaining)
    yield
    for session_id, _name in list(manager.tasks):
        manager.cancel_session_tasks(session_id)


app = FastAPI(title="QuizForge API", version="1.0.0", lifespan=lifespan)
app.add_middleware(CORSMiddleware, allow_origins=CORS_ORIGINS, allow_origin_regex=CORS_ORIGIN_REGEX, allow_credentials=True, allow_methods=["*"], allow_headers=["*"])


@app.get("/api/health")
def health():
    return {"status": "ok"}


@app.post("/api/auth/register", status_code=201)
def register(body: RegisterIn, db: Session = Depends(get_db)):
    if not ALLOW_PUBLIC_REGISTRATION:
        raise HTTPException(403, "Instructor registration is disabled")
    if db.scalar(select(User).where(User.email == body.email.lower())):
        raise HTTPException(409, "Email already registered")
    user = User(name=body.name.strip(), email=body.email.lower(), password_hash=hash_password(body.password))
    db.add(user)
    db.commit()
    db.refresh(user)
    return {"access_token": create_token(user), "user": {"id": user.id, "name": user.name, "email": user.email}}


@app.post("/api/auth/login")
def login(body: LoginIn, db: Session = Depends(get_db)):
    user = db.scalar(select(User).where(User.email == body.email.lower()))
    if not user or not verify_password(body.password, user.password_hash):
        raise HTTPException(401, "Incorrect email or password")
    return {"access_token": create_token(user), "user": {"id": user.id, "name": user.name, "email": user.email}}


@app.post("/api/auth/forgot-password")
def forgot_password(body: ForgotPasswordIn, background_tasks: BackgroundTasks, db: Session = Depends(get_db)):
    if not PASSWORD_RESET_ENABLED:
        raise HTTPException(503, "Password reset email is not configured. Contact your administrator.")
    user = db.scalar(select(User).where(User.email == body.email.lower()))
    if user:
        token = create_password_reset_token(user, PASSWORD_RESET_EXPIRE_MINUTES)
        reset_url = f"{PASSWORD_RESET_BASE_URL}/reset-password?token={token}"
        background_tasks.add_task(send_password_reset_email, user.email, reset_url)
    return {"message": "If that instructor account exists, a reset link has been sent."}


@app.post("/api/auth/reset-password")
def reset_password(body: ResetPasswordIn, db: Session = Depends(get_db)):
    user_id, expected_fingerprint = decode_password_reset_token(body.token)
    user = db.get(User, user_id)
    if not user or not secrets.compare_digest(password_fingerprint(user.password_hash), expected_fingerprint):
        raise HTTPException(400, "This password reset link is invalid or has already been used")
    user.password_hash = hash_password(body.password)
    db.commit()
    return {"message": "Password updated"}


def serialize_quiz(quiz: Quiz, include_answers: bool = True, viewer_id: int | None = None) -> dict:
    questions = []
    for q in quiz.questions:
        item = {"id": q.id, "position": q.position, "text": q.text, "type": q.type, "options": q.options, "match_options": q.match_options or [], "time_limit_sec": q.time_limit_sec, "points": q.points}
        if include_answers:
            item["correct_options"] = q.correct_options
        questions.append(item)
    return {
        "id": quiz.id,
        "title": quiz.title,
        "course_tag": quiz.course_tag,
        "created_at": quiz.created_at,
        "questions": questions,
        "owner": {"id": quiz.owner.id, "name": quiz.owner.name, "email": quiz.owner.email},
        "folder": {"id": quiz.folder.id, "name": quiz.folder.name, "course_tag": quiz.folder.course_tag} if quiz.folder else None,
        "can_edit": viewer_id == quiz.owner_id if viewer_id else False,
    }


def owned_quiz(db: Session, quiz_id: int, user_id: int) -> Quiz:
    quiz = db.scalar(
        select(Quiz)
        .where(Quiz.id == quiz_id, Quiz.owner_id == user_id, Quiz.deleted_at.is_(None))
        .options(selectinload(Quiz.questions), selectinload(Quiz.owner), selectinload(Quiz.folder))
    )
    if not quiz:
        raise HTTPException(404, "Quiz not found")
    return quiz


def accessible_folder(db: Session, folder_id: int, user_id: int, owner_only: bool = False) -> CourseFolder:
    query = (
        select(CourseFolder)
        .where(CourseFolder.id == folder_id)
        .options(
            selectinload(CourseFolder.owner),
            selectinload(CourseFolder.members).selectinload(FolderMember.user),
            selectinload(CourseFolder.quizzes),
        )
    )
    if owner_only:
        query = query.where(CourseFolder.owner_id == user_id)
    else:
        member_folders = select(FolderMember.folder_id).where(FolderMember.user_id == user_id)
        query = query.where(or_(CourseFolder.owner_id == user_id, CourseFolder.id.in_(member_folders)))
    folder = db.scalar(query)
    if not folder:
        raise HTTPException(404, "Course folder not found")
    return folder


def accessible_quiz(db: Session, quiz_id: int, user_id: int) -> Quiz:
    member_folders = select(FolderMember.folder_id).where(FolderMember.user_id == user_id)
    quiz = db.scalar(
        select(Quiz)
        .where(
            Quiz.id == quiz_id,
            Quiz.deleted_at.is_(None),
            or_(
                Quiz.owner_id == user_id,
                Quiz.folder.has(CourseFolder.owner_id == user_id),
                Quiz.folder_id.in_(member_folders),
            ),
        )
        .options(selectinload(Quiz.questions), selectinload(Quiz.owner), selectinload(Quiz.folder))
    )
    if not quiz:
        raise HTTPException(404, "Quiz not found")
    return quiz


def serialize_folder(folder: CourseFolder, viewer_id: int) -> dict:
    return {
        "id": folder.id,
        "name": folder.name,
        "course_tag": folder.course_tag,
        "description": folder.description,
        "created_at": folder.created_at,
        "owner": {"id": folder.owner.id, "name": folder.owner.name, "email": folder.owner.email},
        "is_owner": folder.owner_id == viewer_id,
        "quiz_count": sum(1 for quiz in folder.quizzes if quiz.deleted_at is None),
        "members": [{"id": membership.user.id, "name": membership.user.name, "email": membership.user.email} for membership in folder.members],
    }


def validate_quiz(body: QuizIn):
    for index, question in enumerate(body.questions):
        valid = set(range(len(question.options)))
        if not set(question.correct_options).issubset(valid):
            raise HTTPException(422, f"Question {index + 1} has an invalid correct option")


@app.get("/api/instructors")
def list_instructors(search: str = "", user: User = Depends(current_user), db: Session = Depends(get_db)):
    query = select(User).where(User.id != user.id)
    if search.strip():
        term = f"%{search.strip().lower()}%"
        query = query.where(or_(func.lower(User.name).like(term), func.lower(User.email).like(term)))
    instructors = db.scalars(query.order_by(User.name).limit(20)).all()
    return [{"id": instructor.id, "name": instructor.name, "email": instructor.email} for instructor in instructors]


@app.get("/api/folders")
def list_folders(user: User = Depends(current_user), db: Session = Depends(get_db)):
    member_folders = select(FolderMember.folder_id).where(FolderMember.user_id == user.id)
    folders = db.scalars(
        select(CourseFolder)
        .where(or_(CourseFolder.owner_id == user.id, CourseFolder.id.in_(member_folders)))
        .options(
            selectinload(CourseFolder.owner),
            selectinload(CourseFolder.members).selectinload(FolderMember.user),
            selectinload(CourseFolder.quizzes),
        )
        .order_by(CourseFolder.name)
    ).all()
    return [serialize_folder(folder, user.id) for folder in folders]


@app.post("/api/folders", status_code=201)
def create_folder(body: FolderIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    folder = CourseFolder(owner_id=user.id, name=body.name.strip(), course_tag=body.course_tag.strip(), description=body.description.strip())
    db.add(folder)
    db.commit()
    return serialize_folder(accessible_folder(db, folder.id, user.id), user.id)


@app.put("/api/folders/{folder_id}")
def update_folder(folder_id: int, body: FolderIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    folder = accessible_folder(db, folder_id, user.id, owner_only=True)
    folder.name, folder.course_tag, folder.description = body.name.strip(), body.course_tag.strip(), body.description.strip()
    db.commit()
    return serialize_folder(accessible_folder(db, folder.id, user.id), user.id)


@app.delete("/api/folders/{folder_id}", status_code=204)
def delete_folder(folder_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    folder = accessible_folder(db, folder_id, user.id, owner_only=True)
    for quiz in folder.quizzes:
        quiz.folder_id = None
    db.flush()
    db.delete(folder)
    db.commit()


@app.post("/api/folders/{folder_id}/members", status_code=201)
def add_folder_member(folder_id: int, body: FolderMemberIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    folder = accessible_folder(db, folder_id, user.id, owner_only=True)
    member = db.scalar(select(User).where(User.email == body.email.lower()))
    if not member:
        raise HTTPException(404, "Instructor must create an account before being invited")
    if member.id == user.id:
        raise HTTPException(409, "Folder owner already has access")
    if db.get(FolderMember, (folder.id, member.id)):
        raise HTTPException(409, "Instructor already has access")
    db.add(FolderMember(folder_id=folder.id, user_id=member.id))
    db.commit()
    return serialize_folder(accessible_folder(db, folder.id, user.id), user.id)


@app.delete("/api/folders/{folder_id}/members/{member_id}", status_code=204)
def remove_folder_member(folder_id: int, member_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    folder = accessible_folder(db, folder_id, user.id, owner_only=True)
    membership = db.get(FolderMember, (folder.id, member_id))
    if not membership:
        raise HTTPException(404, "Folder member not found")
    db.delete(membership)
    db.commit()


@app.get("/api/quizzes")
def list_quizzes(search: str = "", user: User = Depends(current_user), db: Session = Depends(get_db)):
    member_folders = select(FolderMember.folder_id).where(FolderMember.user_id == user.id)
    query = (
        select(Quiz)
        .join(Quiz.owner)
        .where(Quiz.deleted_at.is_(None), or_(Quiz.owner_id == user.id, Quiz.folder.has(CourseFolder.owner_id == user.id), Quiz.folder_id.in_(member_folders)))
        .options(selectinload(Quiz.questions), selectinload(Quiz.owner), selectinload(Quiz.folder))
    )
    if search.strip():
        term = f"%{search.strip().lower()}%"
        query = query.where(or_(func.lower(Quiz.title).like(term), func.lower(Quiz.course_tag).like(term), func.lower(User.name).like(term)))
    quizzes = db.scalars(query.order_by(Quiz.created_at.desc())).unique().all()
    return [serialize_quiz(quiz, viewer_id=user.id) for quiz in quizzes]


@app.post("/api/quizzes", status_code=201)
def create_quiz(body: QuizIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    validate_quiz(body)
    if body.folder_id is not None:
        accessible_folder(db, body.folder_id, user.id)
    quiz = Quiz(owner_id=user.id, folder_id=body.folder_id, title=body.title.strip(), course_tag=body.course_tag.strip())
    quiz.questions = [Question(position=i, **q.model_dump()) for i, q in enumerate(body.questions)]
    db.add(quiz)
    db.commit()
    return serialize_quiz(owned_quiz(db, quiz.id, user.id), viewer_id=user.id)


@app.get("/api/quizzes/{quiz_id}")
def get_quiz(quiz_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    return serialize_quiz(accessible_quiz(db, quiz_id, user.id), viewer_id=user.id)


@app.put("/api/quizzes/{quiz_id}")
def update_quiz(quiz_id: int, body: QuizIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    validate_quiz(body)
    quiz = owned_quiz(db, quiz_id, user.id)
    if db.scalar(select(func.count(GameSession.id)).where(GameSession.quiz_id == quiz.id, GameSession.started_at.is_not(None))):
        raise HTTPException(409, "Played quizzes are locked to preserve result history. Create a new quiz version instead.")
    if body.folder_id is not None:
        accessible_folder(db, body.folder_id, user.id)
    quiz.title, quiz.course_tag, quiz.folder_id = body.title.strip(), body.course_tag.strip(), body.folder_id
    quiz.questions.clear()
    db.flush()
    quiz.questions.extend(Question(position=i, **q.model_dump()) for i, q in enumerate(body.questions))
    db.commit()
    return serialize_quiz(owned_quiz(db, quiz_id, user.id), viewer_id=user.id)


@app.delete("/api/quizzes/{quiz_id}", status_code=204)
def delete_quiz(quiz_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    quiz = owned_quiz(db, quiz_id, user.id)
    pending_sessions = db.scalars(
        select(GameSession).where(GameSession.quiz_id == quiz.id, GameSession.started_at.is_(None))
    ).all()
    for game in pending_sessions:
        manager.cancel_session_tasks(game.id)
        db.delete(game)
    quiz.deleted_at = utcnow()
    db.commit()


def new_pin(db: Session) -> str:
    for _ in range(20):
        pin = str(random.randint(100000, 999999))
        if not db.scalar(select(GameSession).where(GameSession.pin == pin, GameSession.status != "ended")):
            return pin
    raise HTTPException(503, "Could not allocate a game PIN")


def session_query(session_id: int):
    return (
        select(GameSession)
        .where(GameSession.id == session_id)
        .options(
            selectinload(GameSession.quiz).selectinload(Quiz.questions),
            selectinload(GameSession.quiz).selectinload(Quiz.owner),
            selectinload(GameSession.quiz).selectinload(Quiz.folder),
            selectinload(GameSession.participants),
            selectinload(GameSession.host),
        )
    )


def accessible_session(db: Session, session_id: int, user_id: int) -> GameSession:
    game = db.scalar(session_query(session_id))
    if not game or user_id not in {game.host_id or game.quiz.owner_id, game.quiz.owner_id}:
        raise HTTPException(404, "Session not found")
    return game


def hosted_session(db: Session, session_id: int, user_id: int) -> GameSession:
    game = db.scalar(session_query(session_id))
    if not game or (game.host_id or game.quiz.owner_id) != user_id:
        raise HTTPException(404, "Session not found")
    return game


def session_json(game: GameSession) -> dict:
    return {"id": game.id, "quiz_id": game.quiz_id, "host_id": game.host_id, "pin": game.pin, "status": game.status, "current_question_index": game.current_question_index, "is_revealed": game.is_revealed, "started_at": game.started_at, "ended_at": game.ended_at}


def load_session(db: Session, session_id: int) -> GameSession | None:
    return db.scalar(
        select(GameSession)
        .where(GameSession.id == session_id)
        .options(selectinload(GameSession.quiz).selectinload(Quiz.questions), selectinload(GameSession.participants))
    )


def schedule_reveal(session_id: int, question_id: int, delay: float):
    manager.set_task(session_id, "reveal", auto_reveal(session_id, question_id, delay))


def schedule_advance(session_id: int, question_id: int):
    manager.set_task(session_id, "advance", auto_advance(session_id, question_id, RESULT_DISPLAY_SECONDS))


async def auto_reveal(session_id: int, question_id: int, delay: float):
    await asyncio.sleep(max(0, delay))
    async with manager.lock(session_id):
        with SessionLocal() as db:
            game = load_session(db, session_id)
            if not game or game.status != "live" or game.is_revealed or game.current_question_index is None:
                return
            question = game.quiz.questions[game.current_question_index]
            if question.id != question_id:
                return
            await reveal_game(db, game)


async def auto_advance(session_id: int, question_id: int, delay: float):
    await asyncio.sleep(max(0, delay))
    async with manager.lock(session_id):
        with SessionLocal() as db:
            game = load_session(db, session_id)
            if not game or game.status != "live" or not game.is_revealed or game.current_question_index is None:
                return
            if game.quiz.questions[game.current_question_index].id != question_id:
                return
            await advance_game(db, game)


async def begin_question(db: Session, game: GameSession, index: int):
    manager.cancel_task(game.id, "advance")
    now = utcnow()
    question = game.quiz.questions[index]
    game.status = "live"
    game.current_question_index = index
    game.question_started_at = now
    game.question_deadline_at = now + timedelta(seconds=question.time_limit_sec)
    game.is_revealed = False
    db.commit()
    await manager.broadcast(game.id, {"type": "question_start", "question": question_payload(game, question), "question_count": len(game.quiz.questions)})
    schedule_reveal(game.id, question.id, question.time_limit_sec)


async def reveal_game(db: Session, game: GameSession) -> dict:
    if game.is_revealed or game.current_question_index is None:
        question = game.quiz.questions[game.current_question_index or 0]
        return {"reveal": reveal_payload(db, game, question), "leaderboard": leaderboard(db, game.id, update_ranks=False)}
    manager.cancel_task(game.id, "reveal")
    game.is_revealed = True
    db.commit()
    question = game.quiz.questions[game.current_question_index]
    reveal_data = {
        **reveal_payload(db, game, question),
        "next_in_sec": RESULT_DISPLAY_SECONDS,
        "next_at": iso(utcnow() + timedelta(seconds=RESULT_DISPLAY_SECONDS)),
    }
    board = leaderboard(db, game.id)
    score_by_participant = {row["participant_id"]: row["score"] for row in board}
    answers = db.scalars(select(Answer).where(Answer.session_id == game.id, Answer.question_id == question.id)).all()
    answer_by_participant = {answer.participant_id: answer for answer in answers}
    personal_results = {}
    for participant in game.participants:
        answer = answer_by_participant.get(participant.id)
        personal_results[participant.id] = {
            "type": "player_result",
            "question_id": question.id,
            "is_correct": bool(answer and answer.is_correct),
            "points_awarded": answer.points_awarded if answer else 0,
            "selected_options": answer.selected_options if answer else [],
            "score": score_by_participant.get(participant.id, 0),
        }
    await manager.send_participant_payloads(game.id, personal_results)
    await manager.broadcast(game.id, {"type": "question_reveal", **reveal_data})
    await manager.broadcast(game.id, {"type": "leaderboard_update", "leaderboard": board})
    schedule_advance(game.id, question.id)
    return {"reveal": reveal_data, "leaderboard": board}


async def finish_game(db: Session, game: GameSession):
    manager.cancel_session_tasks(game.id)
    game.status, game.ended_at = "ended", utcnow()
    db.commit()
    board = leaderboard(db, game.id)
    await manager.broadcast(game.id, {"type": "session_end", "leaderboard": board})
    return session_json(game)


async def advance_game(db: Session, game: GameSession):
    next_index = (game.current_question_index or 0) + 1
    if next_index >= len(game.quiz.questions):
        return await finish_game(db, game)
    await begin_question(db, game, next_index)
    return session_json(game)


@app.post("/api/quizzes/{quiz_id}/sessions", status_code=201)
def create_session(quiz_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    quiz = accessible_quiz(db, quiz_id, user.id)
    if not quiz.questions:
        raise HTTPException(422, "Add at least one question before launching")
    game = GameSession(quiz_id=quiz.id, host_id=user.id, pin=new_pin(db))
    db.add(game)
    db.commit()
    db.refresh(game)
    return session_json(game)


@app.get("/api/sessions/history")
def session_history(user: User = Depends(current_user), db: Session = Depends(get_db)):
    games = db.scalars(
        select(GameSession)
        .join(Quiz)
        .where(or_(GameSession.host_id == user.id, Quiz.owner_id == user.id), GameSession.started_at.is_not(None))
        .options(
            selectinload(GameSession.quiz).selectinload(Quiz.questions),
            selectinload(GameSession.quiz).selectinload(Quiz.owner),
            selectinload(GameSession.participants),
            selectinload(GameSession.host),
        )
        .order_by(GameSession.started_at.desc())
    ).unique().all()
    return [
        {
            "id": game.id,
            "quiz_id": game.quiz_id,
            "quiz_title": game.quiz.title,
            "course_tag": game.quiz.course_tag,
            "pin": game.pin,
            "status": game.status,
            "started_at": game.started_at,
            "ended_at": game.ended_at,
            "participant_count": len(game.participants),
            "question_count": len(game.quiz.questions),
            "host": {"id": (game.host or game.quiz.owner).id, "name": (game.host or game.quiz.owner).name, "email": (game.host or game.quiz.owner).email},
            "quiz_owner": {"id": game.quiz.owner.id, "name": game.quiz.owner.name},
            "hosted_by_me": (game.host_id or game.quiz.owner_id) == user.id,
            "used_my_quiz": game.quiz.owner_id == user.id and (game.host_id or game.quiz.owner_id) != user.id,
            "can_delete": (game.host_id or game.quiz.owner_id) == user.id,
        }
        for game in games
    ]


@app.get("/api/sessions/{session_id}")
def get_session(session_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    game = accessible_session(db, session_id, user.id)
    return session_snapshot(db, game)


@app.get("/api/sessions/{session_id}/statistics")
def session_statistics(session_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    game = accessible_session(db, session_id, user.id)
    board = leaderboard(db, game.id, update_ranks=False)
    answers = db.scalars(select(Answer).where(Answer.session_id == game.id)).all()
    by_question: dict[int, list[Answer]] = {}
    for answer in answers:
        by_question.setdefault(answer.question_id, []).append(answer)
    question_stats = []
    for question in game.quiz.questions:
        question_answers = by_question.get(question.id, [])
        correct_count = sum(1 for answer in question_answers if answer.is_correct)
        question_stats.append(
            {
                "id": question.id,
                "position": question.position,
                "text": question.text,
                "type": question.type,
                "options": question.options,
                "match_options": question.match_options or [],
                "correct_options": question.correct_options,
                "distribution": answer_distribution(db, game, question),
                "response_count": len(question_answers),
                "correct_count": correct_count,
                "correct_rate": round(correct_count / len(game.participants) * 100, 1) if game.participants else 0,
            }
        )
    total_possible = len(game.participants) * len(game.quiz.questions)
    total_correct = sum(1 for answer in answers if answer.is_correct)
    return {
        "session": {
            "id": game.id,
            "quiz_title": game.quiz.title,
            "course_tag": game.quiz.course_tag,
            "pin": game.pin,
            "status": game.status,
            "started_at": game.started_at,
            "ended_at": game.ended_at,
            "host": {"id": (game.host or game.quiz.owner).id, "name": (game.host or game.quiz.owner).name},
        },
        "summary": {
            "participants": len(game.participants),
            "questions": len(game.quiz.questions),
            "average_accuracy": round(total_correct / total_possible * 100, 1) if total_possible else 0,
            "average_score": round(sum(row["score"] for row in board) / len(board)) if board else 0,
        },
        "leaderboard": board,
        "questions": question_stats,
    }


@app.delete("/api/sessions/{session_id}", status_code=204)
def delete_session(session_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    game = hosted_session(db, session_id, user.id)
    manager.cancel_session_tasks(game.id)
    db.delete(game)
    db.commit()


@app.post("/api/sessions/join")
async def join(body: JoinIn, db: Session = Depends(get_db)):
    game = db.scalar(select(GameSession).where(GameSession.pin == body.pin).options(selectinload(GameSession.quiz).selectinload(Quiz.questions), selectinload(GameSession.participants)))
    if not game or game.status == "ended":
        raise HTTPException(404, "No active session uses that PIN")
    if body.resume_token:
        participant = db.scalar(select(Participant).where(Participant.session_id == game.id, Participant.resume_token == body.resume_token))
        if participant:
            return {"session_id": game.id, "participant_id": participant.id, "resume_token": participant.resume_token, "display_name": participant.display_name}
    if game.status != "pending":
        raise HTTPException(409, "This quiz has already started")
    participant = Participant(session_id=game.id, display_name=body.display_name.strip(), student_id=(body.student_id or "").strip() or None, resume_token=secrets.token_urlsafe(32))
    db.add(participant)
    db.commit()
    db.refresh(participant)
    participants = db.scalars(select(Participant).where(Participant.session_id == game.id).order_by(Participant.joined_at)).all()
    await manager.broadcast(game.id, {"type": "player_joined", "participants": [{"id": p.id, "display_name": p.display_name} for p in participants]})
    return {"session_id": game.id, "participant_id": participant.id, "resume_token": participant.resume_token, "display_name": participant.display_name}


@app.post("/api/sessions/{session_id}/start")
async def start_session(session_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    game = hosted_session(db, session_id, user.id)
    if game.status != "pending":
        raise HTTPException(409, "Session is not in the lobby")
    game.started_at = utcnow()
    await begin_question(db, game, 0)
    return session_json(game)


@app.post("/api/sessions/{session_id}/reveal")
async def reveal(session_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    game = hosted_session(db, session_id, user.id)
    if game.status != "live" or game.current_question_index is None:
        raise HTTPException(409, "No live question")
    async with manager.lock(game.id):
        db.expire_all()
        game = hosted_session(db, session_id, user.id)
        return await reveal_game(db, game)


@app.post("/api/sessions/{session_id}/next")
async def next_question(session_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    game = hosted_session(db, session_id, user.id)
    if game.status != "live" or not game.is_revealed:
        raise HTTPException(409, "Reveal the current answer first")
    async with manager.lock(game.id):
        db.expire_all()
        game = hosted_session(db, session_id, user.id)
        if game.status != "live" or not game.is_revealed:
            raise HTTPException(409, "The session has already advanced")
        return await advance_game(db, game)


@app.post("/api/sessions/{session_id}/end")
async def finish_session(session_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    game = hosted_session(db, session_id, user.id)
    if game.status == "ended":
        return session_json(game)
    async with manager.lock(game.id):
        db.expire_all()
        game = hosted_session(db, session_id, user.id)
        return await finish_game(db, game)


def export_rows(db: Session, game: GameSession):
    board = leaderboard(db, game.id, update_ranks=False)
    rank_by_id = {row["participant_id"]: row for row in board}
    answers = db.scalars(select(Answer).where(Answer.session_id == game.id)).all()
    answer_map = {(a.participant_id, a.question_id): a for a in answers}
    rows = []
    for participant in game.participants:
        marks = []
        correct = 0
        for q in game.quiz.questions:
            answer = answer_map.get((participant.id, q.id))
            value = 1 if answer and answer.is_correct else 0
            marks.append(value)
            correct += value
        score = rank_by_id[participant.id]["score"]
        total = len(game.quiz.questions)
        rows.append([participant.display_name, participant.student_id or "", *marks, correct, round(correct / total * 100, 2) if total else 0, score, rank_by_id[participant.id]["rank"]])
    rows.sort(key=lambda row: row[-1])
    return rows


@app.get("/api/sessions/{session_id}/export")
def export(session_id: int, format: str = Query("xlsx", pattern="^(xlsx|csv)$"), user: User = Depends(current_user), db: Session = Depends(get_db)):
    game = accessible_session(db, session_id, user.id)
    headers = ["Student Name", "Student ID", *[f"Q{i + 1}" for i in range(len(game.quiz.questions))], "Total Correct", "Accuracy %", "Total Score", "Rank"]
    rows = export_rows(db, game)
    stem = f"{game.quiz.course_tag or 'quiz'}-{game.pin}-results"
    if format == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(headers)
        writer.writerows(rows)
        return Response(output.getvalue(), media_type="text/csv", headers={"Content-Disposition": f'attachment; filename="{stem}.csv"'})
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Results"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="16324F")
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(max(len(str(c.value or "")) for c in column) + 2, 32)
    binary = io.BytesIO()
    workbook.save(binary)
    return Response(binary.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{stem}.xlsx"'})


@app.websocket("/ws/sessions/{session_id}")
async def websocket_session(websocket: WebSocket, session_id: int, participant_token: str | None = None, access_token: str | None = None):
    db = SessionLocal()
    client = None
    try:
        game = db.scalar(select(GameSession).where(GameSession.id == session_id).options(selectinload(GameSession.quiz).selectinload(Quiz.questions), selectinload(GameSession.participants)))
        if not game:
            await websocket.close(code=4404)
            return
        participant = None
        role = "player"
        if participant_token:
            participant = db.scalar(select(Participant).where(Participant.session_id == session_id, Participant.resume_token == participant_token))
            if not participant:
                await websocket.close(code=4401)
                return
        elif access_token:
            try:
                user_id = decode_user_id(access_token)
            except HTTPException:
                await websocket.close(code=4401)
                return
            if (game.host_id or game.quiz.owner_id) != user_id:
                await websocket.close(code=4403)
                return
            role = "host"
        else:
            await websocket.close(code=4401)
            return
        client = Client(websocket, role, participant.id if participant else None)
        await manager.connect(session_id, client)
        await manager.send(websocket, session_snapshot(db, game, participant.id if participant else None))
        while True:
            data = await websocket.receive_json()
            if role != "player" or data.get("type") != "answer_submitted":
                continue
            db.expire_all()
            game = db.scalar(select(GameSession).where(GameSession.id == session_id).options(selectinload(GameSession.quiz).selectinload(Quiz.questions)))
            if game.status != "live" or game.is_revealed or game.current_question_index is None:
                await manager.send(websocket, {"type": "answer_rejected", "message": "No question is accepting answers"})
                continue
            question = game.quiz.questions[game.current_question_index]
            if data.get("question_id") != question.id:
                await manager.send(websocket, {"type": "answer_rejected", "message": "That question is no longer active"})
                continue
            now = utcnow()
            if not game.question_started_at or not game.question_deadline_at or now > game.question_deadline_at:
                await manager.send(websocket, {"type": "answer_rejected", "message": "Time is up"})
                continue
            try:
                submitted = [int(i) for i in data.get("selected_options", [])]
            except (TypeError, ValueError):
                submitted = []
            if question.type in {"order", "matching"}:
                selected = submitted
                valid_answer = len(selected) == len(question.options) and sorted(selected) == list(range(len(question.options)))
            else:
                selected = sorted(set(submitted))
                valid_answer = bool(selected) and all(0 <= i < len(question.options) for i in selected) and (question.type != "single" or len(selected) == 1)
            if not valid_answer:
                await manager.send(websocket, {"type": "answer_rejected", "message": "Invalid answer selection"})
                continue
            elapsed_ms = max(0, int((now - game.question_started_at).total_seconds() * 1000))
            is_correct = selected == expected_submission(game, question)
            limit_ms = question.time_limit_sec * 1000
            speed_factor = max(0.5, 1.0 - 0.5 * min(elapsed_ms / limit_ms, 1.0))
            points = round(question.points * speed_factor) if is_correct else 0
            answer = Answer(session_id=session_id, participant_id=participant.id, question_id=question.id, selected_options=selected, is_correct=is_correct, response_time_ms=elapsed_ms, points_awarded=points)
            db.add(answer)
            try:
                db.commit()
            except IntegrityError:
                db.rollback()
                await manager.send(websocket, {"type": "answer_rejected", "message": "Answer already submitted"})
                continue
            score = db.scalar(select(func.coalesce(func.sum(Answer.points_awarded), 0)).where(Answer.participant_id == participant.id)) or 0
            await manager.send(websocket, {"type": "answer_accepted", "question_id": question.id, "score": int(score)})
            count = db.scalar(select(func.count(Answer.id)).where(Answer.session_id == session_id, Answer.question_id == question.id)) or 0
            await manager.broadcast(session_id, {"type": "question_progress", "question_id": question.id, "response_count": count}, role="host")
            participant_count = db.scalar(select(func.count(Participant.id)).where(Participant.session_id == session_id)) or 0
            if participant_count and count >= participant_count:
                schedule_reveal(session_id, question.id, 0)
    except WebSocketDisconnect:
        pass
    finally:
        if client:
            manager.disconnect(session_id, websocket)
        db.close()
